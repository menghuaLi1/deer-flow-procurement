"""Deterministic procurement report generation shared by tools and APIs."""

from __future__ import annotations

import json
import re
from collections.abc import Iterable
from datetime import UTC, datetime
from html import escape
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

SUPPORTED_REPORT_FORMATS = {"docx", "pdf", "json", "md"}
FONT_CANDIDATES = (
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.otf",
    "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
)


def _text(value: Any, fallback: str = "") -> str:
    if value in (None, ""):
        return fallback
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, list):
        return "；".join(filter(None, (_text(item) for item in value)))
    if isinstance(value, dict):
        amount = value.get("amount") or value.get("upper_limit") or value.get("value")
        currency = value.get("currency") or value.get("unit")
        note = value.get("note") or value.get("description")
        if amount or currency or note:
            return " ".join(filter(None, (_text(currency), _text(amount), f"({_text(note)})" if note else "")))
        return "；".join(f"{key}: {_text(item)}" for key, item in value.items() if _text(item))
    return str(value).replace("\n", " ").strip()


def _selected_rows(state: dict[str, Any]) -> list[dict[str, Any]]:
    requirements = {str(row.get("id")): row for row in state.get("requirements") or []}
    candidates = {str(row.get("id")): row for row in state.get("candidates") or []}
    rows: list[dict[str, Any]] = []
    for requirement_id, candidate_id in (state.get("selections") or {}).items():
        requirement = requirements.get(str(requirement_id), {})
        candidate = candidates.get(str(candidate_id), {})
        rows.append({"requirement": requirement, "candidate": candidate})
    if not rows:
        for candidate in state.get("candidates") or []:
            requirement = requirements.get(str(candidate.get("requirement_id")), {})
            rows.append({"requirement": requirement, "candidate": candidate})
    return rows


def _markdown(state: dict[str, Any]) -> str:
    project = state.get("project") or {}
    decision = state.get("decision") or {}
    lines = [
        "# 采购决策报告",
        "",
        f"- 项目名称: {_text(project.get('project_name'), '未填写')}",
        f"- 交付地址: {_text(project.get('delivery_address'), '待确认')}",
        f"- 要求交期: {_text(project.get('due_date'), '待确认')}",
        f"- 预算: {_text(project.get('budget'), '未设置')}",
        f"- 品牌要求: {_text(project.get('brand_requirement'), '允许同等品' if project.get('allow_equivalent') else '待确认')}",
        "",
        "## 材料清单",
        "",
        "| 材料 | 规格 | 数量 | 单位 | 品牌/备注 |",
        "| --- | --- | ---: | --- | --- |",
    ]
    for item in state.get("requirements") or []:
        lines.append(f"| {_text(item.get('material_name'))} | {_text(item.get('specification'))} | {_text(item.get('quantity'))} | {_text(item.get('unit'))} | {_text(item.get('brand_requirement') or item.get('notes'))} |")
    lines.extend(["", "## 供应商分配", "", "| 材料 | 供应商 | 匹配度 | 报价 | 交期 | 风险 |", "| --- | --- | ---: | --- | --- | --- |"])
    for row in _selected_rows(state):
        requirement, candidate = row["requirement"], row["candidate"]
        lines.append(
            f"| {_text(requirement.get('material_name') or candidate.get('material_name'))} | "
            f"{_text(candidate.get('supplier_name'), '待选择')} | {_text(candidate.get('match_score'))} | "
            f"{_text(candidate.get('price_range'), '待询价')} | {_text(candidate.get('lead_time'), '待确认')} | "
            f"{_text(candidate.get('risks'))} |"
        )
    lines.extend(["", "## 证据核验", ""])
    for item in state.get("evidence") or []:
        title = _text(item.get("title") or item.get("evidence_type") or item.get("supplier_name"), "证据")
        status = _text(item.get("manual_conclusion") or item.get("status"), "待确认")
        url = _text(item.get("url"))
        lines.append(f"- [{title}]({url}) - {status}" if url else f"- {title} - {status}")
    if not state.get("evidence"):
        lines.append("- 暂无证据记录")
    lines.extend(["", "## 采购建议", "", _text(decision.get("summary"), "尚未生成采购决策。")])
    if decision.get("recommendation"):
        lines.extend(["", f"推荐方案: {_text(decision.get('recommendation'))}"])
    risks = decision.get("risks") or []
    if risks:
        lines.extend(["", "## 风险与人工确认", ""])
        lines.extend(f"- {_text(item)}" for item in risks)
    return "\n".join(lines) + "\n"


def _set_docx_font(document: Document) -> None:
    styles = document.styles
    styles["Normal"].font.name = "Noto Sans CJK SC"
    styles["Normal"].font.size = Pt(10)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        styles[style_name].font.name = "Noto Sans CJK SC"


def _add_docx_table(document: Document, headers: list[str], rows: Iterable[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value


def _write_docx(path: Path, state: dict[str, Any]) -> None:
    document = Document()
    _set_docx_font(document)
    document.add_heading("采购决策报告", 0)
    project = state.get("project") or {}
    for label, key, fallback in (
        ("项目名称", "project_name", "未填写"),
        ("交付地址", "delivery_address", "待确认"),
        ("要求交期", "due_date", "待确认"),
        ("预算", "budget", "未设置"),
        ("品牌要求", "brand_requirement", "待确认"),
    ):
        document.add_paragraph(f"{label}：{_text(project.get(key), fallback)}")
    document.add_heading("材料清单", 1)
    _add_docx_table(
        document,
        ["材料", "规格", "数量", "单位", "品牌/备注"],
        (
            [
                _text(row.get("material_name")),
                _text(row.get("specification")),
                _text(row.get("quantity")),
                _text(row.get("unit")),
                _text(row.get("brand_requirement") or row.get("notes")),
            ]
            for row in state.get("requirements") or []
        ),
    )
    document.add_heading("供应商分配", 1)
    _add_docx_table(
        document,
        ["材料", "供应商", "匹配度", "报价", "交期", "风险"],
        (
            [
                _text(row["requirement"].get("material_name") or row["candidate"].get("material_name")),
                _text(row["candidate"].get("supplier_name"), "待选择"),
                _text(row["candidate"].get("match_score")),
                _text(row["candidate"].get("price_range"), "待询价"),
                _text(row["candidate"].get("lead_time"), "待确认"),
                _text(row["candidate"].get("risks")),
            ]
            for row in _selected_rows(state)
        ),
    )
    document.add_heading("证据核验", 1)
    for item in state.get("evidence") or []:
        text = f"{_text(item.get('title') or item.get('evidence_type'), '证据')}：{_text(item.get('status'), '待确认')}"
        if item.get("url"):
            text += f"\n来源：{_text(item.get('url'))}"
        if item.get("notes") or item.get("manual_note"):
            text += f"\n说明：{_text(item.get('manual_note') or item.get('notes'))}"
        document.add_paragraph(text)
    document.add_heading("采购建议", 1)
    decision = state.get("decision") or {}
    document.add_paragraph(_text(decision.get("summary"), "尚未生成采购决策。"))
    document.add_heading("风险与人工确认", 1)
    risks = list(decision.get("risks") or []) + list(decision.get("pending_confirmations") or [])
    for risk in risks or ["无已记录风险"]:
        document.add_paragraph(_text(risk), style="List Bullet")
    document.save(path)


def _register_pdf_font() -> str:
    for candidate in FONT_CANDIDATES:
        if Path(candidate).exists():
            try:
                pdfmetrics.registerFont(TTFont("NotoCJK", candidate))
                return "NotoCJK"
            except Exception:
                continue
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    return "STSong-Light"


def _write_pdf(path: Path, state: dict[str, Any]) -> None:
    font = _register_pdf_font()
    styles = getSampleStyleSheet()
    title = ParagraphStyle("CJKTitle", parent=styles["Title"], fontName=font, fontSize=20, leading=26, alignment=TA_CENTER)
    heading = ParagraphStyle("CJKHeading", parent=styles["Heading2"], fontName=font, fontSize=13, leading=18, spaceBefore=8)
    body = ParagraphStyle("CJKBody", parent=styles["BodyText"], fontName=font, fontSize=9, leading=14)
    document = SimpleDocTemplate(str(path), pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
    story: list[Any] = [Paragraph("采购决策报告", title), Spacer(1, 5 * mm)]
    project = state.get("project") or {}
    project_rows = [
        ["项目名称", _text(project.get("project_name"), "未填写")],
        ["交付地址", _text(project.get("delivery_address"), "待确认")],
        ["要求交期", _text(project.get("due_date"), "待确认")],
        ["预算", _text(project.get("budget"), "未设置")],
        ["品牌要求", _text(project.get("brand_requirement"), "待确认")],
    ]
    project_table = Table([[Paragraph(escape(cell), body) for cell in row] for row in project_rows], colWidths=[28 * mm, 220 * mm])
    project_table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d8d8de")), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f4f3f7")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("PADDING", (0, 0), (-1, -1), 5)]))
    story.extend([project_table, Paragraph("材料清单", heading)])

    requirement_rows = [["材料", "规格", "数量", "单位", "品牌/备注"]]
    requirement_rows.extend(
        [_text(row.get("material_name")), _text(row.get("specification")), _text(row.get("quantity")), _text(row.get("unit")), _text(row.get("brand_requirement") or row.get("notes"))] for row in state.get("requirements") or []
    )
    req_table = Table([[Paragraph(escape(cell), body) for cell in row] for row in requirement_rows], repeatRows=1, colWidths=[45 * mm, 82 * mm, 24 * mm, 20 * mm, 76 * mm])
    req_table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d8d8de")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#ede9fe")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("PADDING", (0, 0), (-1, -1), 4)]))
    story.extend([req_table, PageBreak(), Paragraph("供应商分配与证据", heading)])
    supplier_rows = [["材料", "供应商", "匹配", "报价", "交期", "风险"]]
    for row in _selected_rows(state):
        requirement, candidate = row["requirement"], row["candidate"]
        supplier_rows.append(
            [
                _text(requirement.get("material_name") or candidate.get("material_name")),
                _text(candidate.get("supplier_name"), "待选择"),
                _text(candidate.get("match_score")),
                _text(candidate.get("price_range"), "待询价"),
                _text(candidate.get("lead_time"), "待确认"),
                _text(candidate.get("risks")),
            ]
        )
    supplier_table = Table([[Paragraph(escape(cell), body) for cell in row] for row in supplier_rows], repeatRows=1, colWidths=[42 * mm, 58 * mm, 20 * mm, 42 * mm, 35 * mm, 50 * mm])
    supplier_table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d8d8de")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dcfce7")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("PADDING", (0, 0), (-1, -1), 4)]))
    story.append(supplier_table)
    for item in state.get("evidence") or []:
        evidence = " · ".join(filter(None, (_text(item.get("title") or item.get("evidence_type"), "证据"), _text(item.get("status"), "待确认"), _text(item.get("url")), _text(item.get("manual_note") or item.get("notes")))))
        story.append(Paragraph(f"• {escape(evidence)}", body))
    story.extend([Paragraph("采购建议与风险", heading), Paragraph(escape(_text((state.get("decision") or {}).get("summary"), "尚未生成采购决策。")), body)])
    for item in (state.get("decision") or {}).get("risks") or []:
        story.append(Paragraph(f"• {escape(_text(item))}", body))
    document.build(story)


def generate_procurement_reports(
    state: dict[str, Any],
    outputs_dir: Path,
    formats: Iterable[str] = ("docx", "pdf", "json"),
    filename_prefix: str = "procurement-decision",
) -> list[Path]:
    requested = list(dict.fromkeys(format.lower() for format in formats))
    unsupported = set(requested) - SUPPORTED_REPORT_FORMATS
    if unsupported:
        raise ValueError(f"Unsupported report format: {', '.join(sorted(unsupported))}")
    if not requested:
        raise ValueError("At least one report format is required")

    outputs_dir.mkdir(parents=True, exist_ok=True)
    safe_prefix = re.sub(r"[^A-Za-z0-9._-]+", "-", filename_prefix.strip() or "procurement-decision").strip("-")
    timestamp = datetime.now(UTC).strftime("%Y%m%d-%H%M%S")
    paths: list[Path] = []
    for format in requested:
        path = outputs_dir / f"{safe_prefix}-{timestamp}.{format}"
        if format == "docx":
            _write_docx(path, state)
        elif format == "pdf":
            _write_pdf(path, state)
        elif format == "json":
            path.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
        else:
            path.write_text(_markdown(state), encoding="utf-8")
        paths.append(path)
    return paths
